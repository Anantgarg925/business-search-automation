from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from config import JV_TARGET_MIN_SCORE
from utils import clean_phone_number


HEADERS = [
    "Rank",
    "Mill Name",
    "City",
    "State",
    "Phone",
    "WhatsApp",
    "Website",
    "Google Rating",
    "Reviews Count",
    "JV Score",
    "Address",
    "Google Maps",
    "Notes",
]


def _to_output_rows(df: pd.DataFrame) -> list[list]:
    rows = []
    for idx, (_, row) in enumerate(df.iterrows(), start=1):
        phone = clean_phone_number(str(row.get("phone_number", "")))
        wa_link = f"https://wa.me/91{phone}" if phone else ""
        maps_link = row.get("maps_url") or (f"https://www.google.com/maps/place/?q=place_id:{row.get('place_id', '')}" if row.get("place_id") else "")
        rows.append([
            idx,
            row.get("name", ""),
            row.get("city", ""),
            row.get("state", ""),
            phone,
            wa_link,
            row.get("website", ""),
            row.get("rating", 0),
            row.get("user_ratings_total", 0),
            row.get("jv_score", 0),
            row.get("formatted_address", ""),
            maps_link,
            "",
        ])
    return rows


def _row_to_output(rank: int, row: pd.Series) -> list:
    phone = clean_phone_number(str(row.get("phone_number", "")))
    wa_link = f"https://wa.me/91{phone}" if phone else ""
    maps_link = row.get("maps_url") or (f"https://www.google.com/maps/place/?q=place_id:{row.get('place_id', '')}" if row.get("place_id") else "")
    return [
        rank,
        row.get("name", ""),
        row.get("city", ""),
        row.get("state", ""),
        phone,
        wa_link,
        row.get("website", ""),
        row.get("rating", 0),
        row.get("user_ratings_total", 0),
        row.get("jv_score", 0),
        row.get("formatted_address", ""),
        maps_link,
        "",
    ]


def _style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(12, max_len + 2), 55)


def _populate_sheet(ws, df: pd.DataFrame, color_scores: bool = False):
    ws.append(HEADERS)
    for row in _to_output_rows(df):
        ws.append(row)

    _style_sheet(ws)

    if color_scores:
        green = PatternFill("solid", fgColor="C6EFCE")
        yellow = PatternFill("solid", fgColor="FFEB9C")
        orange = PatternFill("solid", fgColor="FCE4D6")
        for r in range(2, ws.max_row + 1):
            score_cell = ws[f"J{r}"]
            score = int(score_cell.value or 0)
            if score >= 80:
                score_cell.fill = green
            elif score >= 60:
                score_cell.fill = yellow
            elif score >= 50:
                score_cell.fill = orange


def _summary_sheet(ws, df: pd.DataFrame):
    total = len(df)
    already = (df["classification"] == "ALREADY EXPORTING - SKIP").sum()
    domestic = (df["classification"] == "DOMESTIC ONLY - JV TARGET ✅").sum()
    verify = (df["classification"] == "VERIFY MANUALLY ⚠️").sum()

    ws.append(["Metric", "Value"])
    ws.append(["Total mills found on Google Maps", total])
    ws.append(["Total already exporting (skip)", int(already)])
    ws.append(["Total domestic only (JV targets)", int(domestic)])
    ws.append(["Total to verify manually", int(verify)])
    ws.append([])
    ws.append(["Breakdown by city", "Domestic Targets"])

    city_domestic = (
        df[df["classification"] == "DOMESTIC ONLY - JV TARGET ✅"]
        .groupby("city")
        .size()
        .sort_values(ascending=False)
    )

    for city, count in city_domestic.items():
        ws.append([city, int(count)])

    ws.append([])
    ws.append(["Top 3 recommended cities for field visit", ""])
    for city in city_domestic.head(3).index.tolist():
        ws.append([city, "High domestic target density"])

    _style_sheet(ws)


def generate_excel(df: pd.DataFrame, output_file: str):
    domestic = df[df["classification"] == "DOMESTIC ONLY - JV TARGET ✅"].copy()
    top_targets = domestic[domestic["jv_score"] >= JV_TARGET_MIN_SCORE].sort_values("jv_score", ascending=False)
    verify = df[df["classification"] == "VERIFY MANUALLY ⚠️"].copy()
    exporting = df[df["classification"] == "ALREADY EXPORTING - SKIP"].copy()

    wb = Workbook()
    wb.remove(wb.active)

    ws_top = wb.create_sheet("🎯 TOP JV TARGETS")
    _populate_sheet(ws_top, top_targets, color_scores=True)

    ws_all_domestic = wb.create_sheet("📋 ALL DOMESTIC MILLS")
    _populate_sheet(ws_all_domestic, domestic)

    ws_verify = wb.create_sheet("⚠️ VERIFY MANUALLY")
    ws_verify.append(HEADERS + ["Matched APEDA Name", "Match Score"])
    for idx, (_, row) in enumerate(verify.iterrows(), start=1):
        ws_verify.append(_row_to_output(idx, row) + [row.get("matched_apeda_name", ""), row.get("match_score", 0)])
    _style_sheet(ws_verify)

    ws_exporting = wb.create_sheet("❌ ALREADY EXPORTING")
    ws_exporting.append(HEADERS + ["Match Confidence", "Matched APEDA Name"])
    for idx, (_, row) in enumerate(exporting.iterrows(), start=1):
        ws_exporting.append(_row_to_output(idx, row) + [row.get("match_score", 0), row.get("matched_apeda_name", "")])
    _style_sheet(ws_exporting)

    ws_summary = wb.create_sheet("📊 Summary Stats")
    _summary_sheet(ws_summary, df)

    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
